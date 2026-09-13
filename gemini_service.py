"""
Gemini Service for JobFit AI.

Handles resume and job description analysis using Google's official Gemini SDK.
Produces structured JSON adhering to predefined scoring criteria and schema.
"""

import json
import logging
import os
from pathlib import Path
import re
import time
from typing import Any, Dict, List, Optional
from dotenv import load_dotenv
from pydantic import BaseModel, Field

# Configure logging
logger = logging.getLogger(__name__)

# Ensure .env is loaded
BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")


# ==========================================
# Custom Exception Hierarchy
# ==========================================
class GeminiServiceError(Exception):
    """Base exception for Gemini service errors."""
    pass


class GeminiConfigError(GeminiServiceError):
    """Raised when API key or configuration is missing or invalid."""
    pass


class GeminiInputError(GeminiServiceError):
    """Raised when input (JD text or resume file) is invalid."""
    pass


class GeminiAPIError(GeminiServiceError):
    """Raised when Gemini API request fails (network, rate limit, quota, etc.)."""
    pass


class GeminiParsingError(GeminiServiceError):
    """Raised when Gemini returns malformed, empty, or schema-noncompliant JSON."""
    pass


from scoring_engine import (
    CandidateEvaluationComponents,
    calculate_ats_score,
    calculate_job_match,
)


# ==========================================
# Pydantic Schema for Structured Output
# ==========================================
class MissingSkill(BaseModel):
    skill: str
    importance: str = Field(description="Must be one of: 'critical', 'important', 'preferred'")


class SkillBreakdownItem(BaseModel):
    skill: str
    match_percentage: int = Field(ge=0, le=100)


class RecommendedLearningItem(BaseModel):
    skill: str
    reason: str
    topics: List[str]
    resource_url: Optional[str] = Field(
        default=None,
        description="Authoritative, valid documentation or learning portal URL (e.g. official docs, MDN, roadmap.sh)",
    )


class RoadmapStep(BaseModel):
    step_number: int = Field(description="Step index (1, 2, 3, etc.)")
    phase: str = Field(description="e.g. 'Phase 1: Foundations', 'Phase 2: Core Practical', 'Phase 3: Advanced'")
    skill_focus: str = Field(description="Target skill or competency to acquire")
    duration: str = Field(description="Estimated timeframe, e.g. '1-2 Weeks'")
    key_topics: List[str] = Field(description="Key practical topics to learn")
    learning_resource_url: Optional[str] = Field(
        default=None,
        description="Official documentation or high-quality learning URL (e.g. https://git-scm.com/doc, https://docs.python.org, https://react.dev, https://docs.docker.com, https://roadmap.sh)",
    )


class ComponentScores(BaseModel):
    required_skills: Optional[float] = Field(default=0.0, ge=0, le=100)
    relevant_experience: Optional[float] = Field(default=0.0, ge=0, le=100)
    relevant_keywords: Optional[float] = Field(default=0.0, ge=0, le=100)
    projects: Optional[float] = Field(default=0.0, ge=0, le=100)
    education: Optional[float] = Field(default=0.0, ge=0, le=100)
    resume_structure: Optional[float] = Field(default=0.0, ge=0, le=100)
    preferred_skills: Optional[float] = Field(default=None, description="0-100 score or None if no preferred skills in JD")


class ResumeAnalysisOutput(BaseModel):
    candidate_name: str
    ats_score: Optional[int] = Field(default=0, ge=0, le=100)
    job_match_percentage: Optional[int] = Field(default=0, ge=0, le=100)
    matching_skills: List[str]
    missing_skills: List[MissingSkill]
    partial_skills: List[str]
    skill_breakdown: List[SkillBreakdownItem]
    recommended_learning: List[RecommendedLearningItem]
    learning_roadmap: Optional[List[RoadmapStep]] = Field(
        default_factory=list,
        description="Phased, step-by-step career and skill-gap closure roadmap with durations and authoritative URLs",
    )
    recommendation: str
    component_scores: Optional[ComponentScores] = None


# ==========================================
# System Prompt & Evaluation Criteria
# ==========================================
SYSTEM_INSTRUCTION = """
You are an expert ATS resume evaluator and technical recruiter.
Your task is to analyze a candidate's resume (PDF or DOCX document) against a specific Job Description.

EVALUATION RULES:
1. Analyze ONLY information present in the Job Description and the candidate's resume.
2. NEVER invent candidate skills, experiences, certifications, education, or projects.
3. If an item is not mentioned or unavailable in the resume, mark it as unavailable or missing; do not guess.
4. Distinguish missing skills (explicitly required by the JD but absent in the resume) from unmentioned skills.

SCORING CRITERIA:
Calculate scores strictly using the following predefined weights:

ATS SCORE (0 - 100):
- Required skills match: 40%
- Relevant experience: 20%
- Resume keywords alignment: 15%
- Relevant projects: 10%
- Education relevance: 5%
- Resume structure & ATS readability: 10%

JOB MATCH PERCENTAGE (0 - 100):
- Required core skills: 50%
- Relevant experience: 25%
- Projects: 15%
- Education: 5%
- Preferred skills: 5%

MISSING SKILLS IMPORTANCE:
- 'critical': Fundamental skills required to perform the daily job.
- 'important': Major tools or frameworks mentioned in requirements.
- 'preferred': Nice-to-have or bonus qualifications.

RECOMMENDED LEARNING & ACTIONABLE ROADMAP:
1. Recommend practical skills and topics that would bridge the candidate's skill gaps for this specific job.
2. Provide a structured, phased learning roadmap (e.g. Phase 1: Foundations, Phase 2: Core Practical, Phase 3: Advanced/Production) with estimated durations.
3. For each recommended skill and roadmap step, include real, authoritative documentation/learning URLs (e.g. official docs like https://docs.python.org, https://react.dev, https://git-scm.com/doc, https://docs.docker.com, https://kubernetes.io/docs, https://developer.mozilla.org, https://roadmap.sh, etc.).

Return ONLY valid JSON complying with the requested schema.
"""


def extract_text_from_docx(file_path: str | Path) -> str:
    """
    Extracts text cleanly from a .docx file including paragraphs and tables.

    Args:
        file_path: Path to the .docx file.

    Returns:
        Consolidated plain text string with preserved line breaks.

    Raises:
        GeminiInputError: If the file cannot be opened or contains no readable text.
    """
    path_obj = Path(file_path).resolve()
    if not path_obj.exists():
        raise GeminiInputError(f"DOCX file does not exist: {path_obj}")

    try:
        import docx
    except ImportError as err:
        raise GeminiServiceError(f"python-docx is not installed: {err}")

    try:
        doc = docx.Document(path_obj)
    except Exception as err:
        logger.error(f"Failed to parse DOCX file {path_obj}: {err}")
        raise GeminiInputError(f"Corrupted or unreadable DOCX file: {err}")

    lines = []
    # Extract paragraphs
    for para in doc.paragraphs:
        txt = para.text.strip()
        if txt:
            lines.append(txt)

    # Extract tables
    for table in doc.tables:
        for row in table.rows:
            row_items = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_items:
                # Deduplicate identical adjacent cells (often caused by merged cells)
                deduped = []
                for item in row_items:
                    if not deduped or item != deduped[-1]:
                        deduped.append(item)
                if deduped:
                    lines.append(" | ".join(deduped))

    full_text = "\n".join(lines).strip()
    if not full_text:
        raise GeminiInputError(f"DOCX file '{path_obj.name}' contains no readable text.")

    return full_text


def extract_text_from_pdf(file_path: str | Path) -> str:
    """
    Extracts readable text from a PDF file using pypdf.

    Args:
        file_path: Path to the PDF file.

    Returns:
        Consolidated plain text string.

    Raises:
        GeminiInputError: If file cannot be read or contains no readable text.
    """
    path_obj = Path(file_path).resolve()
    if not path_obj.exists():
        raise GeminiInputError(f"PDF file does not exist: {path_obj}")

    try:
        from pypdf import PdfReader
    except ImportError as err:
        raise GeminiServiceError(f"pypdf is not installed: {err}")

    try:
        reader = PdfReader(str(path_obj))
        pages_text = []
        for page in reader.pages:
            txt = page.extract_text()
            if txt and txt.strip():
                pages_text.append(txt.strip())
        full_text = "\n\n".join(pages_text).strip()
    except Exception as err:
        logger.error(f"Failed extracting text from PDF {path_obj}: {err}")
        raise GeminiInputError(f"Corrupted or unreadable PDF: {err}")

    if not full_text:
        raise GeminiInputError(f"PDF file '{path_obj.name}' contains no readable text.")

    return full_text


def classify_document_content(text: str, filename: str) -> str:
    """
    Classifies a document as 'resume' or 'job_description' based on content and filename heuristics.

    Returns:
        'resume' or 'job_description'
    """
    text_lower = text.lower()
    fname_lower = filename.lower()

    # Strong filename cues
    resume_fname_patterns = [
        r"(^|[_\-\s])(resume|cv|curriculum|biodata|profile|candidate)([_\-\s\d\.]|$)",
    ]
    jd_fname_patterns = [
        r"(^|[_\-\s])(jd|job[_\-\s]*desc|job[_\-\s]*description|job|requirement|position|role|opening|hiring|vacancy)([_\-\s\d\.]|$)",
        r"(^|[_\-\s])jd\d*([_\-\s\.]|$)",
        r"\bjd\b",
    ]

    resume_score = 0
    jd_score = 0

    for pat in jd_fname_patterns:
        if re.search(pat, fname_lower):
            jd_score += 8
            break

    for pat in resume_fname_patterns:
        if re.search(pat, fname_lower):
            resume_score += 8
            break

    # Content cues for Resume:
    # 1. Personal contact indicators
    has_email = bool(re.search(r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+", text))
    has_phone = bool(re.search(r"(\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}", text))
    has_profile_link = "linkedin.com" in text_lower or "github.com" in text_lower

    if has_email:
        resume_score += 4
    if has_phone:
        resume_score += 3
    if has_profile_link:
        resume_score += 3

    has_any_contact = has_email or has_phone or has_profile_link

    # 2. Typical resume sections and career keywords
    resume_keywords = [
        "work experience", "professional experience", "employment history",
        "career objective", "academic projects", "personal projects",
        "cgpa", "gpa", "curriculum vitae", "declaration", "hobbies", "languages known",
        "personal summary", "about me",
    ]
    for kw in resume_keywords:
        if kw in text_lower:
            resume_score += 2

    # Content cues for Job Description:
    jd_keywords = [
        "position summary", "job summary", "role summary",
        "is looking for", "looking for a", "seeking a", "we are looking", "we are seeking", "we are hiring",
        "key responsibilities", "responsibilities", "duties", "what you'll do", "what you will do",
        "required skills", "skills required", "key skills", "must have", "must possess",
        "preferred skills", "desired skills", "nice to have",
        "requirements", "basic qualifications", "preferred qualifications", "minimum qualifications",
        "job description", "years of experience", "experience required", "required experience",
        "about the role", "about the job", "about the company", "who we are", "about us",
        "company:", "location:", "department:", "employment type:",
        "equal opportunity employer", "salary range", "compensation", "perks", "benefits",
    ]
    for kw in jd_keywords:
        if kw in text_lower:
            jd_score += 2

    if jd_score > resume_score:
        return "job_description"
    elif resume_score > jd_score:
        return "resume"
    else:
        # Tied scores: absence of contact info with JD cues heavily favors JD
        if not has_any_contact and jd_score > 0:
            return "job_description"
        if has_any_contact:
            return "resume"
        for pat in jd_fname_patterns:
            if re.search(pat, fname_lower):
                return "job_description"
        for pat in resume_fname_patterns:
            if re.search(pat, fname_lower):
                return "resume"
        return "job_description"


def is_valid_job_description_text(raw_text: str) -> bool:
    """
    Validates whether a text message qualifies as a Job Description.
    Rejects casual conversation, questions, empty inputs, or resume text.

    Returns:
        True if the text has typical JD terminology and length, False otherwise.
    """
    if not raw_text or not isinstance(raw_text, str):
        return False

    cleaned = raw_text.strip()
    if len(cleaned) < 30 or len(cleaned) > 10000:
        return False

    # Check if content has explicit resume cues (e.g. email, phone, personal links)
    has_contact = (
        bool(re.search(r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+", cleaned))
        or bool(re.search(r"(\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}", cleaned))
        or "linkedin.com" in cleaned.lower()
    )
    if has_contact and classify_document_content(cleaned, "text_input.txt") == "resume":
        return False

    cleaned_lower = cleaned.lower()

    jd_keywords = [
        r"\brequirements?\b",
        r"\bqualifications?\b",
        r"\bresponsibilit(y|ies)\b",
        r"\bskills?\b",
        r"\bexperience\b",
        r"\blooking for\b",
        r"\bseeking\b",
        r"\bjob description\b",
        r"\brole\b",
        r"\bposition\b",
        r"\bdeveloper\b",
        r"\bengineer\b",
        r"\bdesigner\b",
        r"\bmanager\b",
        r"\banalyst\b",
        r"\barchitect\b",
        r"\bcandidate\b",
        r"\btechnolog(y|ies)\b",
        r"\bframeworks?\b",
        r"\bproficien(t|cy)\b",
        r"\bhands-on\b",
        r"\bmust have\b",
        r"\bnice to have\b",
        r"\bwhat you('ll| will) do\b",
        r"\babout the (role|job)\b",
        r"\bduties\b",
        r"\bdegree\b",
        r"\bbachelor\b",
        r"\bmaster\b",
        r"\byears?\b",
        r"\bstack\b",
        r"\bprogramming\b",
        r"\bpython\b",
        r"\bjava\b",
        r"\bjavascript\b",
        r"\breact\b",
        r"\bsql\b",
        r"\bcloud\b",
        r"\baws\b",
        r"\bgcp\b",
        r"\bazure\b",
        r"\bdocker\b",
        r"\bkubernetes\b",
        r"\bperks\b",
        r"\bbenefits\b",
        r"\bequal opportunity\b",
        r"\bcollaborat(e|ion)\b",
        r"\bscientist\b",
        r"\bdata\b",
        r"\bmachine learning\b",
        r"\bsoftware\b",
        r"\blead\b",
        r"\bbackend\b",
        r"\bfrontend\b",
        r"\bfullstack\b",
        r"\bdevops\b",
        r"\bsre\b",
        r"\bqa\b",
        r"\btesting\b",
        r"\bapis?\b",
        r"\bpipeline\b",
    ]

    # Structural hiring/role indicators that must be present in a genuine JD
    hiring_indicators = [
        r"\brequirements?\b",
        r"\bqualifications?\b",
        r"\bresponsibilit(y|ies)\b",
        r"\blooking for\b",
        r"\bseeking\b",
        r"\bjob description\b",
        r"\brole\b",
        r"\bposition\b",
        r"\bopening\b",
        r"\bhiring\b",
        r"\bdeveloper\b",
        r"\bengineer\b",
        r"\bscientist\b",
        r"\banalyst\b",
        r"\barchitect\b",
        r"\bmanager\b",
        r"\bexperience\b",
        r"\byears?\b",
        r"\bmust have\b",
        r"\bwhat you('ll| will) do\b",
        r"\bduties\b",
    ]

    has_hiring_indicator = any(re.search(pat, cleaned_lower) for pat in hiring_indicators)
    if not has_hiring_indicator:
        return False

    matched_keywords = 0
    for pat in jd_keywords:
        if re.search(pat, cleaned_lower):
            matched_keywords += 1

    return matched_keywords >= 2


def analyze_resume_against_jd(
    resume_path: str | Path,
    job_description: str,
    api_key: Optional[str] = None,
    model_name: str = "gemini-3.6-flash",
) -> Dict[str, Any]:
    """
    Analyzes a single candidate resume against a Job Description using Gemini.

    Args:
        resume_path: Absolute or relative path to the candidate's PDF or DOCX file.
        job_description: Text of the job description.
        api_key: Optional Gemini API key. If not provided, loaded from environment/GEMINI_API_KEY.
        model_name: Gemini model name (default: gemini-3.6-flash).

    Returns:
        Structured dictionary matching ResumeAnalysisOutput schema.

    Raises:
        GeminiConfigError: If GEMINI_API_KEY is missing or placeholder.
        GeminiInputError: If JD is empty or resume file is missing/corrupted.
        GeminiAPIError: If the Gemini API request encounters network, quota, or auth issues.
        GeminiParsingError: If Gemini returns malformed or non-compliant output.
    """
    # 1. Validate API Key
    effective_key = api_key or os.getenv("GEMINI_API_KEY")
    if not effective_key or effective_key.strip() in ("", "your_gemini_api_key_here"):
        raise GeminiConfigError(
            "GEMINI_API_KEY is missing or not configured in .env file.\n"
            "Please provide a valid API key from https://aistudio.google.com/."
        )

    # 2. Validate Job Description
    if not job_description or not job_description.strip():
        raise GeminiInputError("Job Description cannot be empty or whitespace-only.")

    # 3. Validate Resume File
    path_obj = Path(resume_path).resolve()
    if not path_obj.exists():
        raise GeminiInputError(f"Resume file does not exist at: {path_obj}")

    if not path_obj.is_file():
        raise GeminiInputError(f"Resume path is not a file: {path_obj}")

    if path_obj.stat().st_size == 0:
        raise GeminiInputError(f"Resume file is empty (0 bytes): {path_obj}")

    suffix = path_obj.suffix.lower()
    if suffix not in (".pdf", ".docx"):
        raise GeminiInputError(
            f"Unsupported resume file type: '{suffix}'. JobFit AI supports .pdf and .docx resumes."
        )

    is_docx = (suffix == ".docx")

    try:
        file_bytes = path_obj.read_bytes()
    except OSError as err:
        raise GeminiInputError(f"Failed to read resume file: {err}")

    if is_docx:
        if not file_bytes.startswith(b"PK\x03\x04"):
            raise GeminiInputError("Invalid DOCX file: missing ZIP header magic bytes.")
        docx_text = extract_text_from_docx(path_obj)
    else:
        if not file_bytes.startswith(b"%PDF"):
            raise GeminiInputError("Invalid PDF file: missing %PDF header magic bytes.")

    # 4. Import Google GenAI SDK
    try:
        from google import genai
        from google.genai import types
        from google.genai.errors import APIError
    except ImportError as err:
        raise GeminiServiceError(f"google-genai SDK is not installed: {err}")

    # 5. Prepare Gemini Client and Request
    try:
        client = genai.Client(api_key=effective_key.strip())

        if is_docx:
            prompt_text = (
                f"JOB DESCRIPTION:\n{job_description.strip()}\n\n"
                f"CANDIDATE RESUME TEXT ({path_obj.name}):\n{docx_text}\n\n"
                "Analyze the candidate resume above against this Job Description.\n"
                "Extract candidate skills, calculate ATS score and Job Match % based on the specified criteria, "
                "identify matching, missing, and partial skills, and provide recommended learning."
            )
            contents = [prompt_text]
        else:
            prompt_text = (
                f"JOB DESCRIPTION:\n{job_description.strip()}\n\n"
                "Analyze the attached candidate resume against this Job Description.\n"
                "Extract candidate skills, calculate ATS score and Job Match % based on the specified criteria, "
                "identify matching, missing, and partial skills, and provide recommended learning."
            )
            pdf_part = types.Part.from_bytes(data=file_bytes, mime_type="application/pdf")
            contents = [pdf_part, prompt_text]

        config = types.GenerateContentConfig(
            system_instruction=SYSTEM_INSTRUCTION,
            response_mime_type="application/json",
            response_schema=ResumeAnalysisOutput,
            temperature=0.2,
        )

        MAX_RETRIES = 3
        response = None
        for attempt in range(1, MAX_RETRIES + 1):
            try:
                logger.info(
                    f"Submitting resume '{path_obj.name}' ({suffix}) to Gemini ({model_name}) [Attempt {attempt}/{MAX_RETRIES}]..."
                )
                response = client.models.generate_content(
                    model=model_name,
                    contents=contents,
                    config=config,
                )
                break
            except APIError as api_err:
                err_code = getattr(api_err, "code", None)
                err_msg = str(api_err).lower()
                is_transient = (
                    err_code in (429, 500, 503)
                    or "503" in err_msg
                    or "429" in err_msg
                    or "high demand" in err_msg
                    or "temporarily unavailable" in err_msg
                    or "rate limit" in err_msg
                )
                if is_transient and attempt < MAX_RETRIES:
                    wait_sec = 2.5 * attempt
                    logger.warning(
                        f"Gemini API transient error on attempt {attempt}: {api_err}. Retrying in {wait_sec}s..."
                    )
                    time.sleep(wait_sec)
                else:
                    logger.error(f"Gemini API error during candidate analysis: {api_err}")
                    raise GeminiAPIError(f"Gemini API call failed: {api_err.message or api_err}")

    except GeminiServiceError:
        raise
    except Exception as exc:
        err_str = str(exc)
        if "API_KEY_INVALID" in err_str or "unauthorized" in err_str.lower() or "401" in err_str:
            raise GeminiConfigError(f"Invalid GEMINI_API_KEY: {err_str}")
        logger.error(f"Unexpected error calling Gemini API: {exc}")
        raise GeminiAPIError(f"Unexpected error communicating with Gemini API: {exc}")

    # 6. Validate and Parse Response
    if not response or not response.text:
        raise GeminiParsingError("Gemini returned an empty response.")

    try:
        raw_json = json.loads(response.text)
    except json.JSONDecodeError as json_err:
        resp_len = len(response.text) if response.text else 0
        logger.error(f"Failed to decode Gemini JSON response (length: {resp_len} chars): {json_err}")
        raise GeminiParsingError(f"Gemini returned invalid JSON: {json_err}")

    # Validate against Pydantic schema
    try:
        validated = ResumeAnalysisOutput.model_validate(raw_json)
        result_dict = validated.model_dump()

        # Deterministically compute final scores using Python's scoring engine
        if result_dict.get("component_scores"):
            cs = result_dict["component_scores"]
            result_dict["ats_score"] = calculate_ats_score(
                required_skills=cs.get("required_skills"),
                relevant_experience=cs.get("relevant_experience"),
                relevant_keywords=cs.get("relevant_keywords"),
                projects=cs.get("projects"),
                education=cs.get("education"),
                resume_structure=cs.get("resume_structure"),
            )
            result_dict["job_match_percentage"] = calculate_job_match(
                required_skills=cs.get("required_skills"),
                relevant_experience=cs.get("relevant_experience"),
                projects=cs.get("projects"),
                education=cs.get("education"),
                preferred_skills=cs.get("preferred_skills"),
            )
        else:
            # Fallback: ensure existing scores are strictly clamped
            result_dict["ats_score"] = max(0, min(100, int(result_dict.get("ats_score", 0))))
            result_dict["job_match_percentage"] = max(0, min(100, int(result_dict.get("job_match_percentage", 0))))

        return result_dict
    except Exception as val_err:
        logger.error(f"Gemini JSON did not match expected schema: {val_err}")
        raise GeminiParsingError(f"Response schema validation failed: {val_err}")
