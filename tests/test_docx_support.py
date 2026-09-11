"""
Tests for DOCX support: Job Description and Resume uploads, extraction, and evaluation.
"""

import os
from pathlib import Path
import shutil
from unittest.mock import AsyncMock, MagicMock, patch
import docx
import pytest
from telegram.constants import ParseMode

from bot import (
    analyze_command,
    handle_document_before_jd,
    handle_resume_document,
)
from gemini_service import (
    GeminiInputError,
    analyze_resume_against_jd,
    extract_text_from_docx,
)
from session_manager import (
    clear_session,
    get_all_sessions,
    get_job_description,
    get_resumes,
    get_session,
    set_job_description,
)


@pytest.fixture(autouse=True)
def clean_environment(tmp_path):
    """Clean sessions and temp files before each test."""
    get_all_sessions().clear()
    temp_root = Path(__file__).resolve().parent.parent / "temp" / "users"
    if temp_root.exists():
        shutil.rmtree(temp_root, ignore_errors=True)
    yield
    if temp_root.exists():
        shutil.rmtree(temp_root, ignore_errors=True)


def create_sample_docx(
    file_path: Path,
    paragraphs: list[str] | None = None,
    table_data: list[list[str]] | None = None,
) -> Path:
    """Helper to generate a real, valid .docx file."""
    doc = docx.Document()
    if paragraphs:
        for p in paragraphs:
            doc.add_paragraph(p)
    if table_data:
        table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
        for r_idx, row in enumerate(table_data):
            for c_idx, cell_value in enumerate(row):
                table.cell(r_idx, c_idx).text = cell_value
    doc.save(str(file_path))
    return file_path


def make_mock_docx_update(
    user_id: int,
    file_path: Path,
    file_name: str | None = None,
    mime_type: str = "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
):
    """Creates a mock Telegram update representing a document upload pointing to file_path."""
    update = MagicMock()
    user = MagicMock()
    user.id = user_id
    user.first_name = "Charlie"
    update.effective_user = user

    document = MagicMock()
    document.file_name = file_name or file_path.name
    document.mime_type = mime_type
    document.file_size = file_path.stat().st_size if file_path.exists() else 1024
    document.file_id = f"tg_docx_{user_id}"

    message = MagicMock()
    message.document = document
    message.reply_text = AsyncMock()
    update.message = message

    context = MagicMock()
    bot = MagicMock()

    async def mock_download_to_drive(custom_path):
        os.makedirs(os.path.dirname(custom_path), exist_ok=True)
        shutil.copyfile(file_path, custom_path)

    mock_file = MagicMock()
    mock_file.download_to_drive = AsyncMock(side_effect=mock_download_to_drive)
    bot.get_file = AsyncMock(return_value=mock_file)
    context.bot = bot

    return update, context


# ==========================================
# 1. Tests for extract_text_from_docx
# ==========================================

def test_extract_text_from_valid_docx(tmp_path):
    """Test extracting paragraphs and table content from a valid .docx."""
    docx_path = tmp_path / "sample.docx"
    create_sample_docx(
        docx_path,
        paragraphs=[
            "Senior Backend Engineer",
            "Must have 5+ years of experience in Python, FastAPI, and Docker.",
        ],
        table_data=[
            ["Skill", "Proficiency"],
            ["Python", "Expert"],
            ["PostgreSQL", "Advanced"],
        ],
    )

    extracted = extract_text_from_docx(docx_path)
    assert "Senior Backend Engineer" in extracted
    assert "FastAPI" in extracted
    assert "Skill | Proficiency" in extracted
    assert "Python | Expert" in extracted


def test_extract_text_from_empty_docx(tmp_path):
    """Test that extracting text from a DOCX with no text raises GeminiInputError."""
    empty_docx = tmp_path / "empty.docx"
    create_sample_docx(empty_docx, paragraphs=["   \n  "], table_data=[])

    with pytest.raises(GeminiInputError, match="contains no readable text"):
        extract_text_from_docx(empty_docx)


def test_extract_text_from_nonexistent_file(tmp_path):
    """Test that extract_text_from_docx raises GeminiInputError on missing file."""
    with pytest.raises(GeminiInputError, match="does not exist"):
        extract_text_from_docx(tmp_path / "missing.docx")


def test_extract_text_from_corrupted_docx(tmp_path):
    """Test that a non-ZIP file renamed to .docx raises GeminiInputError."""
    corrupt_docx = tmp_path / "corrupt.docx"
    corrupt_docx.write_bytes(b"This is definitely not a docx file format")

    with pytest.raises(GeminiInputError, match="Corrupted or unreadable DOCX"):
        extract_text_from_docx(corrupt_docx)


# ==========================================
# 2. Tests for Job Description via DOCX
# ==========================================

@pytest.mark.asyncio
async def test_submit_job_description_via_docx(tmp_path):
    """Test user uploading a .docx file before JD is set stores JD and advances state."""
    user_id = 901
    jd_docx = tmp_path / "Job_Description.docx"
    create_sample_docx(
        jd_docx,
        paragraphs=[
            "We are seeking a Lead Data Scientist with 7+ years of experience in Machine Learning, "
            "Python, PyTorch, SQL, and cloud infrastructure (GCP/AWS). Candidate must have strong "
            "communication skills and leadership experience."
        ],
    )

    update, context = make_mock_docx_update(user_id=user_id, file_path=jd_docx)
    await handle_resume_document(update, context)

    # Verify session state and stored JD
    session = get_session(user_id)
    assert session["job_description"] is not None
    assert "Lead Data Scientist" in session["job_description"]
    assert "Machine Learning" in session["job_description"]
    assert session["state"] == "waiting_for_resumes"

    update.message.reply_text.assert_awaited_once()
    reply, _ = update.message.reply_text.call_args
    assert "Job Description received" in reply[0]
    assert "Job_Description.docx" in reply[0]


@pytest.mark.asyncio
async def test_submit_short_job_description_via_docx(tmp_path):
    """Test uploading a .docx JD with fewer than 30 characters is rejected."""
    user_id = 902
    jd_docx = tmp_path / "Too_Short.docx"
    create_sample_docx(jd_docx, paragraphs=["Python dev"])

    update, context = make_mock_docx_update(user_id=user_id, file_path=jd_docx)
    await handle_resume_document(update, context)

    session = get_session(user_id)
    assert session["job_description"] is None
    reply, _ = update.message.reply_text.call_args
    assert "too short" in reply[0]


# ==========================================
# 3. Tests for Resume Upload via DOCX
# ==========================================

@pytest.mark.asyncio
async def test_upload_docx_resume(tmp_path):
    """Test uploading a .docx resume after setting JD."""
    user_id = 903
    set_job_description(user_id, "Looking for Senior Django Developer with AWS experience.")

    resume_docx = tmp_path / "John_Doe_Resume.docx"
    create_sample_docx(
        resume_docx,
        paragraphs=[
            "John Doe - Senior Software Engineer",
            "Experience: 6 years building Django and Python microservices on AWS EC2/S3.",
            "Education: BS in Computer Science.",
        ],
    )

    update, context = make_mock_docx_update(user_id=user_id, file_path=resume_docx)
    await handle_resume_document(update, context)

    resumes = get_resumes(user_id)
    assert len(resumes) == 1
    assert resumes[0]["original_filename"] == "John_Doe_Resume.docx"
    assert resumes[0]["file_type"] == "docx"
    assert Path(resumes[0]["local_path"]).exists()

    reply, _ = update.message.reply_text.call_args
    assert "Resume 1 received" in reply[0]
    assert "John_Doe_Resume.docx" in reply[0]


@pytest.mark.asyncio
async def test_mixed_pdf_and_docx_resumes(tmp_path):
    """Test user uploading a PDF resume and a DOCX resume in the same session."""
    user_id = 904
    set_job_description(user_id, "Looking for Senior Django Developer.")

    # 1. Upload PDF
    pdf_path = tmp_path / "candidate1.pdf"
    pdf_path.write_bytes(b"%PDF-1.4 mock pdf content for candidate 1")
    update1, context1 = make_mock_docx_update(user_id=user_id, file_path=pdf_path, file_name="candidate1.pdf", mime_type="application/pdf")
    await handle_resume_document(update1, context1)

    # 2. Upload DOCX
    docx_path = tmp_path / "candidate2.docx"
    create_sample_docx(docx_path, paragraphs=["Candidate 2 resume content with Python and Django."])
    update2, context2 = make_mock_docx_update(user_id=user_id, file_path=docx_path, file_name="candidate2.docx")
    await handle_resume_document(update2, context2)

    resumes = get_resumes(user_id)
    assert len(resumes) == 2
    assert resumes[0]["file_type"] == "pdf"
    assert resumes[1]["file_type"] == "docx"


# ==========================================
# 4. Tests for Gemini Service with DOCX
# ==========================================

def test_analyze_resume_against_jd_with_docx(tmp_path):
    """Test that analyze_resume_against_jd extracts DOCX text and calls Gemini correctly."""
    resume_docx = tmp_path / "Candidate.docx"
    create_sample_docx(
        resume_docx,
        paragraphs=[
            "Alice Smith - Cloud Engineer",
            "Skills: Python, Kubernetes, Terraform, Docker, AWS.",
        ],
    )

    mock_gemini_response = MagicMock()
    mock_gemini_response.text = (
        '{"candidate_name": "Alice Smith", "ats_score": 85, "job_match_percentage": 90, '
        '"matching_skills": ["Python", "Kubernetes"], "missing_skills": [], '
        '"partial_skills": [], "skill_breakdown": [], "recommended_learning": [], '
        '"recommendation": "Strong candidate for Cloud Engineer role.", '
        '"component_scores": {"required_skills": 90, "relevant_experience": 85, "relevant_keywords": 80, '
        '"projects": 80, "education": 80, "resume_structure": 90, "preferred_skills": 85}}'
    )

    with patch("google.genai.Client") as mock_client_cls:
        mock_client = MagicMock()
        mock_client.models.generate_content.return_value = mock_gemini_response
        mock_client_cls.return_value = mock_gemini_client = mock_client

        result = analyze_resume_against_jd(
            resume_path=resume_docx,
            job_description="Looking for Cloud Engineer with Python, Kubernetes, Terraform.",
            api_key="mock_valid_key",
        )

        assert result["candidate_name"] == "Alice Smith"
        assert result["job_match_percentage"] > 0
        assert result["ats_score"] > 0

        # Verify prompt passed text rather than PDF Part
        call_kwargs = mock_gemini_client.models.generate_content.call_args.kwargs
        contents = call_kwargs["contents"]
        assert len(contents) == 1
        assert "CANDIDATE RESUME TEXT" in contents[0]
        assert "Alice Smith" in contents[0]
